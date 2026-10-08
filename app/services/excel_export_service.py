import io
from datetime import date, datetime
from typing import Any, Dict, List, Optional
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ExcelExportService:
    """
    Standardized, professional Excel (.xlsx) workbook generator for CloudKey PMS reports.
    Produces presentation-grade spreadsheets with branded header, structured metadata,
    styled table headers, cell formatting (currency, dates, percentages), and totals.
    """

    def generate_report_workbook(
        self,
        report_title: str,
        property_name: str,
        date_range_label: str,
        columns: List[Dict[str, Any]],
        rows: List[Dict[str, Any]],
        summary_totals: Optional[Dict[str, Any]] = None,
        sheet_name: str = "Report Data",
    ) -> io.BytesIO:
        """
        Builds an Excel workbook in-memory and returns a BytesIO buffer.

        :param report_title: E.g., 'Arrivals and Departures Report'
        :param property_name: E.g., 'Grand Palace Resort'
        :param date_range_label: E.g., 'Date: 2026-10-08' or 'From: 2026-10-01 To: 2026-10-08'
        :param columns: List of dicts: [
            {'key': 'booking_number', 'header': 'Booking #', 'width': 16, 'align': 'left'},
            {'key': 'amount', 'header': 'Total (₹)', 'width': 14, 'align': 'right', 'format': '#,##0.00'},
            ...
        ]
        :param rows: List of row data dicts matching column keys
        :param summary_totals: Optional dict mapping column keys to total values
        :param sheet_name: Title of the primary sheet
        """
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]  # Excel limits sheet title to 31 chars

        # Fonts & Styles
        title_font = Font(name="Calibri", size=15, bold=True, color="1E293B")
        subtitle_font = Font(name="Calibri", size=11, bold=True, color="475569")
        meta_font = Font(name="Calibri", size=10, italic=True, color="64748B")
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        data_font = Font(name="Calibri", size=10, color="0F172A")
        total_font = Font(name="Calibri", size=10, bold=True, color="0F172A")

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

        thin_border_side = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=thin_border_side,
        )

        total_border = Border(
            top=Side(border_style="thin", color="0F172A"),
            bottom=Side(border_style="double", color="0F172A"),
            left=thin_border_side,
            right=thin_border_side,
        )

        # ── 1. Header Block ──
        ws.append([f"CloudKey PMS — {property_name}"])
        ws.cell(row=1, column=1).font = title_font

        ws.append([report_title])
        ws.cell(row=2, column=1).font = subtitle_font

        generated_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws.append([f"{date_range_label} | Generated on: {generated_str}"])
        ws.cell(row=3, column=1).font = meta_font

        ws.append([])  # Row 4 blank spacing

        # ── 2. Table Column Headers ──
        header_row_idx = 5
        header_values = [col.get("header", col["key"]) for col in columns]
        ws.append(header_values)

        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            align_val = col.get("align", "left")
            cell.alignment = Alignment(horizontal=align_val, vertical="center", wrap_text=True)

        # ── 3. Data Rows ──
        current_row_idx = header_row_idx + 1
        for row_data in rows:
            row_values = []
            for col in columns:
                val = row_data.get(col["key"])
                if isinstance(val, (date, datetime)):
                    val = str(val)
                elif val is None:
                    val = "-"
                row_values.append(val)
            ws.append(row_values)

            # Apply cell styles
            is_even = (current_row_idx % 2 == 0)
            for col_idx, col in enumerate(columns, start=1):
                cell = ws.cell(row=current_row_idx, column=col_idx)
                cell.font = data_font
                cell.border = cell_border
                if is_even:
                    cell.fill = zebra_fill

                align_val = col.get("align", "left")
                cell.alignment = Alignment(horizontal=align_val, vertical="center")

                # Number formatting
                num_format = col.get("format")
                if num_format and cell.value not in (None, "-", ""):
                    try:
                        cell.value = float(cell.value)
                        cell.number_format = num_format
                    except (ValueError, TypeError):
                        pass

            current_row_idx += 1

        # ── 4. Summary / Totals Row ──
        if summary_totals:
            total_values = []
            for col in columns:
                val = summary_totals.get(col["key"], "")
                total_values.append(val)
            ws.append(total_values)

            for col_idx, col in enumerate(columns, start=1):
                cell = ws.cell(row=current_row_idx, column=col_idx)
                cell.font = total_font
                cell.fill = total_fill
                cell.border = total_border
                align_val = col.get("align", "left")
                cell.alignment = Alignment(horizontal=align_val, vertical="center")

                num_format = col.get("format")
                if num_format and cell.value not in (None, "-", ""):
                    try:
                        cell.value = float(cell.value)
                        cell.number_format = num_format
                    except (ValueError, TypeError):
                        pass

        # ── 5. Auto Column Widths ──
        for col_idx, col in enumerate(columns, start=1):
            col_letter = get_column_letter(col_idx)
            configured_width = col.get("width")
            if configured_width:
                ws.column_dimensions[col_letter].width = configured_width
            else:
                header_len = len(str(col.get("header", "")))
                max_len = header_len
                for r in range(header_row_idx + 1, current_row_idx + (1 if summary_totals else 0)):
                    val = ws.cell(row=r, column=col_idx).value
                    if val is not None:
                        max_len = max(max_len, len(str(val)))
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        # Freeze panes at data rows
        ws.freeze_panes = f"A{header_row_idx + 1}"

        # Write to BytesIO
        output_stream = io.BytesIO()
        wb.save(output_stream)
        output_stream.seek(0)
        return output_stream


excel_export_service = ExcelExportService()
